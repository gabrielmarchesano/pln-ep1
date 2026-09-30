from __future__ import annotations

import hashlib
import html
import re
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.model_selection import StratifiedGroupKFold

TEXT_COLUMN = "resp_text"
LABEL_COLUMN = "clarity"
LABELS = ("c1", "c234", "c5")
EXPECTED_LABELS = set(LABELS)


def normalize_text(text: object) -> str:
    """Keep accents, punctuation and negation; merge whitespace and HTML entities."""
    value = "" if pd.isna(text) else str(text)
    return re.sub(r"\s+", " ", html.unescape(value)).strip().lower()


def file_digest(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def read_frame(path: Path, *, training: bool) -> pd.DataFrame:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("O arquivo de dados deve ter extensão .xlsx.")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        if not header or any(name is None for name in header):
            raise ValueError("Cabeçalhos ausentes ou vazios na primeira planilha.")
        if len(set(header)) != len(header):
            raise ValueError("A planilha contém cabeçalhos duplicados.")
        required = {TEXT_COLUMN, LABEL_COLUMN} if training else {TEXT_COLUMN}
        missing = required - set(header)
        if missing:
            raise ValueError(f"Colunas obrigatórias ausentes: {sorted(missing)}")
        frame = pd.DataFrame(list(rows), columns=header)
        if frame.empty:
            raise ValueError("A planilha não contém exemplos.")
        column = list(header).index(TEXT_COLUMN) + 1
        if any(row[0].data_type == "f" for row in sheet.iter_rows(
            min_row=2, min_col=column, max_col=column
        )):
            raise ValueError("resp_text deve conter texto, não fórmulas Excel.")
    finally:
        workbook.close()
    invalid_text = frame[TEXT_COLUMN].map(
        lambda value: not pd.isna(value) and not isinstance(value, str)
    )
    # Excel may infer numeric cells in an otherwise textual column. Keep the row
    # and report this conversion; never edit the source workbook.
    frame.attrs["coerced_text_excel_rows"] = (np.flatnonzero(invalid_text) + 2).tolist()
    frame[TEXT_COLUMN] = frame[TEXT_COLUMN].fillna("").astype(str)
    if training:
        labels = set(frame[LABEL_COLUMN].dropna())
        if frame[LABEL_COLUMN].isna().any() or labels != EXPECTED_LABELS:
            raise ValueError(f"O treino deve conter exatamente os rótulos {list(LABELS)}.")
        if frame[TEXT_COLUMN].map(normalize_text).eq("").all():
            raise ValueError("Todos os textos de treinamento estão vazios.")
    return frame


def load_training_data(path: Path) -> pd.DataFrame:
    return read_frame(path, training=True)[[TEXT_COLUMN, LABEL_COLUMN]].copy()


def text_groups(texts) -> np.ndarray:
    return np.asarray([
        hashlib.sha256(normalize_text(text).encode("utf-8")).hexdigest()
        for text in texts
    ])


def describe_dataset(frame: pd.DataFrame) -> dict:
    normalized = frame[TEXT_COLUMN].map(normalize_text)
    counts = Counter(frame[LABEL_COLUMN])
    grouped = frame.assign(_text=normalized).groupby("_text")[LABEL_COLUMN]
    sizes, unique_labels = grouped.size(), grouped.nunique()
    lengths = normalized.str.split().str.len()
    return {
        "rows": len(frame),
        "class_counts": dict(sorted(counts.items())),
        "majority_accuracy": max(counts.values()) / len(frame),
        "empty_text_rows": int(normalized.eq("").sum()),
        "coerced_text_excel_rows": frame.attrs.get("coerced_text_excel_rows", []),
        "unique_normalized_texts": int(normalized.nunique()),
        "duplicate_extra_rows": int(len(frame) - normalized.nunique()),
        "duplicate_groups": int(sizes.gt(1).sum()),
        "conflicting_duplicate_groups": int(unique_labels.gt(1).sum()),
        "rows_in_conflicting_duplicate_groups": int(sizes[unique_labels.gt(1)].sum()),
        "median_words": float(lengths.median()),
        "mean_words": float(lengths.mean()),
        "p95_words": float(lengths.quantile(0.95)),
        "max_words": int(lengths.max()),
    }


def grouped_splits(texts, labels, groups, folds: int, seed: int) -> list:
    if folds < 2:
        raise ValueError("A validação exige ao menos dois folds.")
    for label in LABELS:
        count = len(np.unique(groups[labels == label]))
        if count < folds:
            raise ValueError(f"A classe {label} tem {count} grupos para {folds} folds.")
    splitter = StratifiedGroupKFold(n_splits=folds, shuffle=True, random_state=seed)
    splits = list(splitter.split(texts, labels, groups))
    for train, valid in splits:
        if set(groups[train]) & set(groups[valid]):
            raise ValueError("Vazamento de grupos entre treino e validação.")
        if set(labels[train]) != EXPECTED_LABELS or set(labels[valid]) != EXPECTED_LABELS:
            raise ValueError("Um fold perdeu uma classe; reduza folds ou altere a semente.")
    return splits
