from __future__ import annotations

import os
import tempfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .data import EXPECTED_LABELS, LABEL_COLUMN, read_frame


def require_new_file(path: Path) -> None:
    if path.exists():
        raise ValueError(f"O destino já existe; escolha um novo caminho: {path}")


def write_predictions(test_path: Path, output_path: Path, predictions) -> None:
    """Edit only label cells, preserving order, other sheets and Excel formatting."""
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("A entrega deve ter extensão .xlsx.")
    if test_path.resolve() == output_path.resolve():
        raise ValueError("O destino não pode sobrescrever a planilha de entrada.")
    require_new_file(output_path)
    frame = read_frame(test_path, training=False)
    predictions = np.asarray(predictions)
    if predictions.ndim != 1 or len(predictions) != len(frame):
        raise ValueError("A quantidade de rótulos não corresponde às linhas de teste.")
    if not set(predictions).issubset(EXPECTED_LABELS):
        raise ValueError("O modelo produziu rótulos inválidos.")
    workbook = load_workbook(test_path)
    try:
        sheet = workbook.worksheets[0]
        headers = [cell.value for cell in sheet[1]]
        if LABEL_COLUMN in headers:
            label_column = headers.index(LABEL_COLUMN) + 1
        else:
            # Appending avoids moving existing columns or invalidating formulas.
            label_column = len(headers) + 1
            sheet.cell(row=1, column=label_column, value=LABEL_COLUMN)
        for row, label in enumerate(predictions, 2):
            sheet.cell(row=row, column=label_column, value=str(label))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary = tempfile.mkstemp(suffix=".xlsx", dir=output_path.parent)
        os.close(descriptor)
        try:
            workbook.save(temporary)
            check = read_frame(Path(temporary), training=False)
            if len(check) != len(frame) or check[LABEL_COLUMN].tolist() != predictions.tolist():
                raise ValueError("Falha ao conferir os rótulos gravados.")
            for column in frame.columns:
                if column != LABEL_COLUMN and not check[column].equals(frame[column]):
                    raise ValueError(f"A coluna {column} foi alterada na gravação.")
            os.replace(temporary, output_path)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
    finally:
        workbook.close()

