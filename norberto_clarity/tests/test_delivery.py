import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from clarity.data import file_digest
from clarity.delivery import write_predictions


@pytest.mark.parametrize("existing_labels", [False, True])
def test_submission_preserves_rows_cells_styles_and_sheets(tmp_path, existing_labels):
    source, destination = tmp_path / "test.xlsx", tmp_path / "nested" / "labeled.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "resp_text", "formula"] + (["clarity"] if existing_labels else []))
    for index, text in enumerate(["resposta repetida", None, "resposta repetida", 123], 1):
        sheet.append([index, text, f"=A{index + 1}+1"] + ([None] if existing_labels else []))
    sheet["B2"].font = Font(bold=True)
    sheet.freeze_panes = "B2"
    workbook.create_sheet("Instruções")["A1"] = "Preserve esta aba."
    workbook.save(source)
    digest = file_digest(source)
    labels = ["c1", "c234", "c5", "c1"]
    write_predictions(source, destination, labels)
    original, result = load_workbook(source), load_workbook(destination)
    assert file_digest(source) == digest
    assert original.sheetnames == result.sheetnames
    assert result.active.max_row == original.active.max_row
    assert result.active.max_column == 4
    assert [result.active.cell(row, 4).value for row in range(2, 6)] == labels
    for row in range(1, 6):
        for column in range(1, 4):
            before, after = original.active.cell(row, column), result.active.cell(row, column)
            assert (before.value, before.data_type, before._style) == (after.value, after.data_type, after._style)
    assert result.active.freeze_panes == "B2"
    assert result["Instruções"]["A1"].value == "Preserve esta aba."


def test_invalid_delivery_never_overwrites_source(training_path, tmp_path):
    digest = file_digest(training_path)
    with pytest.raises(ValueError, match="sobrescrever"):
        write_predictions(training_path, training_path, ["c1"])
    with pytest.raises(ValueError, match="quantidade"):
        write_predictions(training_path, tmp_path / "out.xlsx", ["c1"])
    with pytest.raises(ValueError, match="inválidos"):
        write_predictions(training_path, tmp_path / "out.xlsx", ["c6"] * 108)
    with pytest.raises(ValueError, match="extensão"):
        write_predictions(training_path, tmp_path / "out.csv", ["c1"] * 108)
    assert file_digest(training_path) == digest

